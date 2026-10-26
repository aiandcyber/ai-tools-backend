import tempfile
from pathlib import Path

import pymupdf

from app.catalog import BY_ID
from app.tools import ai, images, media, office, pdf, search, utility
from app.workspace import loads_json, zip_paths


class Outcome:
    def __init__(
        self,
        message: str,
        files: list[Path] | None = None,
        results: list[dict] | None = None,
        meta: dict | None = None,
    ):
        self.message = message
        self.files = files or []
        self.results = results or []
        self.meta = meta or {}


def run_tool(tool_id: str, files: list[Path], options: dict) -> Outcome:
    if tool_id not in BY_ID:
        raise KeyError(tool_id)
    work = Path(tempfile.mkdtemp(prefix="aitool-"))
    return _DISPATCH[tool_id](files, options, work)


def _one(files: list[Path]) -> Path:
    if not files:
        raise ValueError("Upload a file")
    return files[0]


def _opt(options: dict, name: str, default: str = "") -> str:
    value = options.get(name, default)
    return "" if value is None else str(value)


def _pdf_merge(files, options, work):
    if len(files) < 2:
        raise ValueError("Upload at least two PDF files")
    dest = pdf.merge_pdfs(files, work / "combined.pdf")
    return Outcome("Combined PDF", [dest])


def _pdf_compress(files, options, work):
    dest = pdf.compress_pdf(_one(files), work / "compressed.pdf")
    return Outcome("Compressed PDF", [dest])


def _pdf_create(files, options, work):
    dest = pdf.create_pdf(work / "document.pdf", _opt(options, "title"), _opt(options, "body"))
    return Outcome("Created PDF", [dest])


def _word_to_pdf(files, options, work):
    dest = pdf.word_to_pdf(_one(files), work)
    return Outcome("Converted Word to PDF", [dest])


def _model_to_pdf(files, options, work):
    dest = pdf.model_to_pdf(_one(files), work / "model.pdf")
    return Outcome("Created a three-view PDF", [dest])


def _pdf_pages(files, options, work):
    outputs = pdf.transform_pages(
        _one(files),
        work / "pages.pdf",
        _opt(options, "action", "split"),
        _opt(options, "pages"),
        int(_opt(options, "angle", "90") or "90"),
    )
    if len(outputs) > 1:
        archive = zip_paths(outputs, work / "pages.zip")
        return Outcome(f"Wrote {len(outputs)} pages", [archive])
    return Outcome("Updated PDF pages", outputs)


def _pdf_protect(files, options, work):
    dest = pdf.protect_pdf(_one(files), work / "protected.pdf", _opt(options, "password"))
    return Outcome("Protected PDF", [dest])


def _pdf_unlock(files, options, work):
    dest = pdf.unlock_pdf(_one(files), work / "unlocked.pdf", _opt(options, "password"))
    return Outcome("Unlocked PDF", [dest])


def _pdf_decorate(files, options, work):
    dest = pdf.decorate_pdf(
        _one(files),
        work / "decorated.pdf",
        _opt(options, "watermark", "CONFIDENTIAL"),
        _opt(options, "page_numbers", "yes") == "yes",
    )
    return Outcome("Added watermark and page numbers", [dest])


def _pdf_extract_text(files, options, work):
    dest = work / "text.txt"
    dest.write_text(pdf.extract_text(_one(files)), encoding="utf-8")
    return Outcome("Extracted text", [dest])


def _pdf_extract_images(files, options, work):
    saved = pdf.extract_images(_one(files), work)
    return Outcome(f"Extracted {len(saved)} images", [zip_paths(saved, work / "images.zip")])


def _images_to_pdf(files, options, work):
    if not files:
        raise ValueError("Upload at least one image")
    return Outcome("Created PDF", [pdf.images_to_pdf(files, work / "images.pdf")])


def _pdf_to_images(files, options, work):
    saved = pdf.pdf_to_images(_one(files), work)
    return Outcome(f"Rendered {len(saved)} pages", [zip_paths(saved, work / "pages.zip")])


def _ocr(files, options, work):
    dest = work / "ocr.txt"
    dest.write_text(pdf.ocr_file(_one(files)), encoding="utf-8")
    return Outcome("OCR complete", [dest])


def _pdf_md(files, options, work):
    dest = work / "document.md"
    dest.write_text(pdf.pdf_to_markdown(_one(files)), encoding="utf-8")
    return Outcome("Wrote Markdown", [dest])


def _docx_text(files, options, work):
    fmt = _opt(options, "format", "text")
    dest = work / ("document.html" if fmt == "html" else "document.txt")
    dest.write_text(office.docx_to_text(_one(files), fmt), encoding="utf-8")
    return Outcome("Converted DOCX", [dest])


def _xlsx_create(files, options, work):
    dest = office.csv_text_to_xlsx(_opt(options, "csv"), _opt(options, "sheet", "Sheet1"), work / "workbook.xlsx")
    return Outcome("Created workbook", [dest])


def _csv_xlsx(files, options, work):
    return Outcome("Converted CSV", [office.csv_file_to_xlsx(_one(files), work / "workbook.xlsx")])


def _xlsx_csv(files, options, work):
    return Outcome("Exported CSV", [office.xlsx_to_csv(_one(files), work / "sheet.csv")])


def _pptx(files, options, work):
    return Outcome("Created presentation", [office.create_pptx(_opt(options, "slides"), work / "slides.pptx")])


def _xlsx_merge(files, options, work):
    if len(files) < 2:
        raise ValueError("Upload at least two workbooks")
    return Outcome("Merged workbooks", [office.merge_xlsx(files, work / "merged.xlsx")])


def _docx_merge(files, options, work):
    if len(files) < 2:
        raise ValueError("Upload at least two Word files")
    return Outcome("Merged Word documents", [office.merge_docx(files, work / "merged.docx")])


def _pptx_merge(files, options, work):
    if len(files) < 2:
        raise ValueError("Upload at least two PowerPoint files")
    return Outcome("Merged presentations", [office.merge_pptx(files, work / "merged.pptx")])


def _image_compress(files, options, work):
    fmt = _opt(options, "format", "jpeg")
    dest = work / f"compressed.{ 'jpg' if fmt == 'jpeg' else 'webp'}"
    images.compress_image(_one(files), dest, int(_opt(options, "quality", "70") or "70"), fmt)
    return Outcome("Compressed image", [dest])


def _image_transform(files, options, work):
    src = _one(files)
    dest = work / f"transformed{src.suffix or '.png'}"
    images.transform_image(src, dest, _opt(options, "action", "resize"), _opt(options, "width"), _opt(options, "height"), int(_opt(options, "angle", "90") or "90"))
    return Outcome("Transformed image", [dest])


def _image_convert(files, options, work):
    fmt = _opt(options, "format", "png")
    ext = "jpg" if fmt == "jpeg" else fmt
    return Outcome("Converted image", [images.convert_image(_one(files), work / f"converted.{ext}", fmt)])


def _image_watermark(files, options, work):
    src = _one(files)
    dest = work / f"watermarked{src.suffix or '.png'}"
    return Outcome("Watermarked image", [images.watermark_image(src, dest, _opt(options, "text", "DRAFT"))])


def _strip(files, options, work):
    src = _one(files)
    dest = work / f"clean{src.suffix or '.png'}"
    return Outcome("Removed EXIF", [images.strip_exif(src, dest)])


def _bg(files, options, work):
    return Outcome("Removed background", [images.remove_background(_one(files), work / "cutout.png")])


def _favicon(files, options, work):
    return Outcome("Created favicon", [images.favicon(_one(files), work / "favicon.ico")])


def _palette(files, options, work):
    dest = work / "palette.csv"
    dest.write_text(images.palette(_one(files)), encoding="utf-8")
    return Outcome("Extracted palette", [dest])


def _youtube_mp4(files, options, work):
    dest = media.download_youtube_mp4(_opt(options, "url"), work, _opt(options, "quality", "720"))
    return Outcome("Downloaded YouTube video", [dest])


def _youtube_mp3(files, options, work):
    dest = media.download_youtube_mp3(_opt(options, "url"), work, _opt(options, "quality", "192"))
    return Outcome("Downloaded YouTube audio", [dest])


def _media_convert(files, options, work):
    fmt = _opt(options, "format", "mp3")
    return Outcome("Converted media", [media.convert_media(_one(files), work / f"converted.{fmt}")])


def _media_trim(files, options, work):
    src = _one(files)
    dest = work / f"trimmed{src.suffix or '.mp4'}"
    return Outcome("Trimmed media", [media.trim_media(src, dest, _opt(options, "start", "0"), _opt(options, "end"))])


def _media_compress(files, options, work):
    dest = media.compress_media(_one(files), work / "compressed", _opt(options, "kind", "video"))
    return Outcome("Compressed media", [dest])


def _extract_audio(files, options, work):
    return Outcome("Extracted audio", [media.extract_audio(_one(files), work / "audio.mp3")])


def _tts(files, options, work):
    text = _opt(options, "text")
    language = _opt(options, "language", "en") or "en"
    lang = language.strip().lower()
    if lang.startswith("zh"):
        lang = "zh"
    detected = media._detect_tts_lang(text, "en")
    spoken = text
    if detected != lang:
        spoken = ai.translate(text, lang, "auto")
        message = f"Translated to {ai.LANG.get(lang, lang)} and created speech"
    else:
        message = "Created speech"
    return Outcome(message, [media.text_to_speech(spoken, lang, work / "speech.mp3")])


def _stt(files, options, work):
    dest = work / "transcript.txt"
    dest.write_text(media.speech_to_text(_one(files)), encoding="utf-8")
    return Outcome("Transcribed audio", [dest])


def _zip_pack(files, options, work):
    if not files:
        raise ValueError("Upload at least one file")
    return Outcome("Created ZIP", [utility.pack_zip(files, work / "archive.zip")])


def _zip_extract(files, options, work):
    return Outcome("Extracted ZIP", [utility.extract_zip(_one(files), work, work / "contents.zip")])


def _data(files, options, work):
    text = _opt(options, "text") or _one(files).read_text(encoding="utf-8")
    target = _opt(options, "target", "json")
    dest = work / f"data.{target if target != 'yaml' else 'yaml'}"
    dest.write_text(utility.convert_data(text, target), encoding="utf-8")
    return Outcome("Converted data", [dest])


def _md_html(files, options, work):
    text = _opt(options, "text") or _one(files).read_text(encoding="utf-8")
    dest = work / "document.html"
    dest.write_text(utility.markdown_to_html(text), encoding="utf-8")
    return Outcome("Rendered HTML", [dest])


def _md_pdf(files, options, work):
    text = _opt(options, "text") or _one(files).read_text(encoding="utf-8")
    dest = work / "document.pdf"
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_textbox(pymupdf.Rect(72, 72, page.rect.width - 72, page.rect.height - 72), text, fontsize=11)
    doc.save(dest)
    doc.close()
    return Outcome("Created PDF", [dest])


def _qr(files, options, work):
    return Outcome("Created QR code", [utility.qr_png(_opt(options, "text"), work / "qr.png")])


def _qr_scan(files, options, work):
    dest = work / "qr.txt"
    dest.write_text(utility.scan_qr(_one(files)), encoding="utf-8")
    return Outcome("Read QR code", [dest])


def _barcode(files, options, work):
    return Outcome("Created barcode", [utility.barcode_png(_opt(options, "text"), work / "barcode.png")])


def _barcode_scan(files, options, work):
    dest = work / "barcode.txt"
    dest.write_text(utility.scan_barcode(_one(files)), encoding="utf-8")
    return Outcome("Read barcode", [dest])


def _hash(files, options, work):
    data = _opt(options, "text").encode() if _opt(options, "text") else _one(files).read_bytes()
    dest = work / "hash.txt"
    dest.write_text(utility.hash_bytes(data, _opt(options, "algorithm", "sha256")), encoding="utf-8")
    return Outcome("Hashed input", [dest])


def _b64(files, options, work):
    action = _opt(options, "action", "encode")
    raw = _opt(options, "text").encode() if _opt(options, "text") else _one(files).read_bytes()
    dest = work / ("encoded.txt" if action == "encode" else "decoded.bin")
    dest.write_bytes(utility.base64_convert(raw, action))
    return Outcome("Converted Base64", [dest])


def _uuid(files, options, work):
    dest = work / "uuids.txt"
    dest.write_text(utility.make_uuids(int(_opt(options, "count", "1") or "1")), encoding="utf-8")
    return Outcome("Generated UUIDs", [dest])


def _diff(files, options, work):
    dest = work / "diff.txt"
    dest.write_text(utility.diff_text(_opt(options, "left"), _opt(options, "right")), encoding="utf-8")
    return Outcome("Compared text", [dest])


def _case(files, options, work):
    dest = work / "case.txt"
    dest.write_text(utility.change_case(_opt(options, "text"), _opt(options, "style", "lower")), encoding="utf-8")
    return Outcome("Converted case", [dest])


def _web_search(files, options, work):
    payload = search.web_search(
        query=_opt(options, "query"),
        mode=_opt(options, "mode", "auto"),
        region=_opt(options, "region"),
        language=_opt(options, "language", "en"),
        interests=_opt(options, "interests", ""),
        industry=_opt(options, "industry", ""),
        location=_opt(options, "location", ""),
        limit=int(_opt(options, "limit", "20") or 20),
    )
    return Outcome(payload["message"], results=payload["results"], meta=payload.get("meta"))


def _text_file(message: str, work: Path, name: str) -> Outcome:
    dest = work / name
    dest.write_text(message, encoding="utf-8")
    return Outcome(message[:500], [dest])


def _summarize(files, options, work):
    text = ai.summarize(ai.source_text(files, _opt(options, "text")))
    dest = work / "summary.md"
    dest.write_text(text, encoding="utf-8")
    return Outcome(text, [dest])


def _translate(files, options, work):
    text = ai.translate(
        ai.source_text(files, _opt(options, "text")),
        _opt(options, "language", "en"),
        _opt(options, "source_language", "auto"),
    )
    dest = work / "translation.md"
    dest.write_text(text, encoding="utf-8")
    return Outcome(text, [dest])


def _rewrite(files, options, work):
    text = ai.rewrite(ai.source_text(files, _opt(options, "text")), _opt(options, "instruction"))
    dest = work / "rewrite.md"
    dest.write_text(text, encoding="utf-8")
    return Outcome(text, [dest])


def _generate(files, options, work):
    fmt = _opt(options, "format", "docx")
    request = _opt(options, "request")
    preset = _opt(options, "template_preset")
    template_hint = ""
    if preset:
        template_hint = f"Well-known template style: {preset}."
    if files:
        try:
            uploaded = ai.source_text(files, "")
            if uploaded.strip():
                template_hint = (
                    f"{template_hint}\nUploaded template / sample structure:\n{uploaded}"
                ).strip()
        except Exception:
            pass
    body = ai.draft(request, fmt, template_hint=template_hint)
    if fmt == "pdf":
        return _pdf_create([], {"title": "Document", "body": body}, work)
    if fmt == "pptx":
        return Outcome("Created presentation", [office.create_pptx(body, work / "slides.pptx")])
    document = work / "document.docx"
    from docx import Document
    doc = Document()
    for paragraph in body.split("\n"):
        doc.add_paragraph(paragraph)
    doc.save(document)
    return Outcome("Created document", [document])


def export_text(text: str, fmt: str, work: Path | None = None) -> Outcome:
    """Export already-generated text to a downloadable format."""
    content = (text or "").strip()
    if not content:
        raise ValueError("Nothing to export")
    fmt = (fmt or "txt").strip().lower()
    root = work or Path(tempfile.mkdtemp(prefix="aitool-export-"))
    if fmt in {"txt", "text"}:
        dest = root / "export.txt"
        dest.write_text(content, encoding="utf-8")
        return Outcome("Exported text", [dest])
    if fmt in {"md", "markdown"}:
        dest = root / "export.md"
        dest.write_text(content, encoding="utf-8")
        return Outcome("Exported Markdown", [dest])
    if fmt == "pdf":
        return _pdf_create([], {"title": "Export", "body": content}, root)
    if fmt == "docx":
        document = root / "export.docx"
        from docx import Document
        doc = Document()
        for paragraph in content.split("\n"):
            doc.add_paragraph(paragraph)
        doc.save(document)
        return Outcome("Exported Word document", [document])
    raise ValueError("Unsupported export format")


def _classify(files, options, work):
    dest = work / "classification.txt"
    dest.write_text(ai.classify(_one(files)), encoding="utf-8")
    return Outcome("Classified file", [dest])


_DISPATCH = {
    "pdf_merge": _pdf_merge,
    "pdf_compress": _pdf_compress,
    "pdf_create": _pdf_create,
    "word_to_pdf": _word_to_pdf,
    "model_to_pdf": _model_to_pdf,
    "pdf_pages": _pdf_pages,
    "pdf_protect": _pdf_protect,
    "pdf_unlock": _pdf_unlock,
    "pdf_decorate": _pdf_decorate,
    "pdf_extract_text": _pdf_extract_text,
    "pdf_extract_images": _pdf_extract_images,
    "images_to_pdf": _images_to_pdf,
    "pdf_to_images": _pdf_to_images,
    "ocr": _ocr,
    "pdf_to_markdown": _pdf_md,
    "docx_to_text": _docx_text,
    "xlsx_create": _xlsx_create,
    "csv_to_xlsx": _csv_xlsx,
    "xlsx_to_csv": _xlsx_csv,
    "pptx_create": _pptx,
    "xlsx_merge": _xlsx_merge,
    "docx_merge": _docx_merge,
    "pptx_merge": _pptx_merge,
    "image_compress": _image_compress,
    "image_transform": _image_transform,
    "image_convert": _image_convert,
    "image_watermark": _image_watermark,
    "image_strip_exif": _strip,
    "image_remove_bg": _bg,
    "favicon": _favicon,
    "palette": _palette,
    # "youtube_mp4": _youtube_mp4,  # off website; use ytube-download app
    # "youtube_mp3": _youtube_mp3,
    "media_convert": _media_convert,
    "media_trim": _media_trim,
    "media_compress": _media_compress,
    "extract_audio": _extract_audio,
    "tts": _tts,
    "stt": _stt,
    "zip_pack": _zip_pack,
    "zip_extract": _zip_extract,
    "data_convert": _data,
    "markdown_html": _md_html,
    "markdown_pdf": _md_pdf,
    "qr_generate": _qr,
    "qr_scan": _qr_scan,
    "barcode_generate": _barcode,
    "barcode_scan": _barcode_scan,
    "hash_text": _hash,
    "base64_tool": _b64,
    "uuid_tool": _uuid,
    "text_diff": _diff,
    "case_convert": _case,
    "web_search": _web_search,
    "summarize": _summarize,
    "translate": _translate,
    "rewrite": _rewrite,
    "generate_document": _generate,
    "classify": _classify,
}

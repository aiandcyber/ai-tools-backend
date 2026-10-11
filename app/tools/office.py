import csv
import io
from pathlib import Path

from docx import Document
from openpyxl import Workbook, load_workbook
from pptx import Presentation


def docx_to_text(src: Path, fmt: str) -> str:
    document = Document(str(src))
    paragraphs = [paragraph.text for paragraph in document.paragraphs]
    if fmt == "html":
        body = "".join(f"<p>{_escape(text)}</p>" for text in paragraphs)
        return f"<!doctype html><html><body>{body}</body></html>"
    return "\n".join(paragraphs)


def _escape(text: str) -> str:
    return text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def csv_text_to_xlsx(text: str, sheet_name: str, dest: Path) -> Path:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = (sheet_name or "Sheet1")[:31]
    for row in csv.reader(io.StringIO(text)):
        sheet.append(row)
    workbook.save(dest)
    return dest


def csv_file_to_xlsx(src: Path, dest: Path) -> Path:
    return csv_text_to_xlsx(src.read_text(encoding="utf-8"), "Sheet1", dest)


def xlsx_to_csv(src: Path, dest: Path) -> Path:
    workbook = load_workbook(src, data_only=True)
    sheet = workbook.active
    with dest.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle)
        for row in sheet.iter_rows(values_only=True):
            writer.writerow(["" if value is None else value for value in row])
    return dest


def create_pptx(text: str, dest: Path) -> Path:
    presentation = Presentation()
    chunks = [chunk.strip() for chunk in text.split("\n---\n")]
    for chunk in chunks or [text]:
        if not chunk:
            continue
        lines = chunk.splitlines()
        layout = presentation.slide_layouts[1]
        slide = presentation.slides.add_slide(layout)
        slide.shapes.title.text = lines[0][:120]
        body = "\n".join(lines[1:]).strip()
        if body and len(slide.placeholders) > 1:
            slide.placeholders[1].text = body
    if not presentation.slides:
        slide = presentation.slides.add_slide(presentation.slide_layouts[0])
        slide.shapes.title.text = "Untitled"
    presentation.save(dest)
    return dest


def merge_xlsx(files: list[Path], dest: Path) -> Path:
    output = Workbook()
    output.remove(output.active)
    for index, item in enumerate(files, start=1):
        source = load_workbook(item, data_only=True)
        origin = source.active
        sheet = output.create_sheet(f"{index}-{origin.title}"[:31])
        for row in origin.iter_rows(values_only=True):
            sheet.append(list(row))
    output.save(dest)
    return dest


def merge_docx(files: list[Path], dest: Path) -> Path:
    """Append DOCX bodies in upload order into one document."""
    from copy import deepcopy

    if len(files) < 2:
        raise ValueError("Upload at least two Word files")
    combined = Document(str(files[0]))
    body = combined.element.body
    for path in files[1:]:
        combined.add_page_break()
        source = Document(str(path))
        for child in list(source.element.body):
            # Keep one sectPr on the combined document
            if child.tag.endswith("sectPr"):
                continue
            body.append(deepcopy(child))
    combined.save(dest)
    return dest


def merge_pptx(files: list[Path], dest: Path) -> Path:
    """Append slides from each presentation into one PPTX (best-effort shape copy)."""
    from copy import deepcopy

    if len(files) < 2:
        raise ValueError("Upload at least two PowerPoint files")
    merged = Presentation(str(files[0]))
    for path in files[1:]:
        source = Presentation(str(path))
        for slide in source.slides:
            layout = merged.slide_layouts[min(6, len(merged.slide_layouts) - 1)]
            new_slide = merged.slides.add_slide(layout)
            for shape in list(new_slide.shapes):
                shape.element.getparent().remove(shape.element)
            for shape in slide.shapes:
                new_slide.shapes._spTree.insert_element_before(deepcopy(shape.element), "p:extLst")
    merged.save(dest)
    return dest
